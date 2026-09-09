# /// script
# requires-python = ">=3.12"
# dependencies = [
#     "openpyxl"
# ]
# ///

import argparse
import os

import openpyxl

parser = argparse.ArgumentParser(description="Analyze combined_verdicts.xlsx for correct results across runs")
parser.add_argument("infile", nargs="?", default="results/combined_verdicts.xlsx", help="Path to the input xlsx file")
parser.add_argument("outfile", nargs="?", default="results/combined_verdicts_analysis.xlsx", help="Path to the output xlsx file")
parser.add_argument("--sheet", default=None, help="Sheet name to analyze (defaults to the active sheet)")
args = parser.parse_args()

CORRECT = "correct"

wb = openpyxl.load_workbook(args.infile)
ws = wb[args.sheet] if args.sheet else wb.active

header = [c.value for c in ws[1]]
runs = [(h[: -len("_result")], i) for i, h in enumerate(header) if isinstance(h, str) and h.endswith("_result")]
if not runs:
    raise SystemExit("Error: no '*_result' columns found; cannot detect runs")

run_names = [name for name, _ in runs]
count_col = len(header) + 1
first_col = len(header) + 2
ws.cell(row=1, column=count_col, value="correct_count")
ws.cell(row=1, column=first_col, value="first_correct_run")

per_run_correct = {name: 0 for name in run_names}
total_rows = 0
any_correct = 0
all_correct = 0
dist = {}

for r in range(2, ws.max_row + 1):
    total_rows += 1
    correct_names = [name for name, idx in runs if str(ws.cell(row=r, column=idx + 1).value).strip().lower() == CORRECT]
    n = len(correct_names)
    for name in correct_names:
        per_run_correct[name] += 1
    ws.cell(row=r, column=count_col, value=n)
    ws.cell(row=r, column=first_col, value=correct_names[0] if correct_names else None)
    dist[n] = dist.get(n, 0) + 1
    if n > 0:
        any_correct += 1
    if n == len(runs):
        all_correct += 1

outdir = os.path.dirname(args.outfile)
if outdir:
    os.makedirs(outdir, exist_ok=True)
wb.save(args.outfile)

print(f"Input:  {args.infile}")
print(f"Output: {args.outfile}")
print(f"Runs detected (column order): {', '.join(run_names)}")
print(f"Total rows: {total_rows}")
print()
print("Correct results per run:")
for name in run_names:
    print(f"  {name}: {per_run_correct[name]}")
print()
print(f"Tasks correct in at least one run: {any_correct}")
print(f"Tasks correct in all {len(runs)} runs: {all_correct}")
print("Distribution by number of correct runs:")
for k in sorted(dist):
    print(f"  {k}/{len(runs)} correct: {dist[k]}")
